"""Convertit les documents de référence en fiches Markdown pour la base privée.

Usage (depuis backend/) :
    python -m scripts.importer_documents                       # tout knowledge/prive/sources/
    python -m scripts.importer_documents chemin/doc.pdf --acces public
    python -m scripts.importer_documents --categorie produit fiche1.pdf fiche2.pdf

Formats : .pdf, .docx, .xlsx, .html/.htm, .md, .txt.
Les fiches produites vont dans knowledge/prive/ (exclu de Git). Par défaut elles
sont « internes » : seul le profil commercial y a accès. Utilisez --acces public
pour des documents diffusables à tous (ex. fiches techniques produits).
Relancez ensuite le serveur (ou POST /api/recharger) pour les prendre en compte.
"""

from __future__ import annotations

import argparse
import re
import sys
import unicodedata
from datetime import date
from pathlib import Path

RACINE = Path(__file__).resolve().parents[2]
DOSSIER_PRIVE = RACINE / "knowledge" / "prive"
DOSSIER_SOURCES = DOSSIER_PRIVE / "sources"
FORMATS = {".pdf", ".docx", ".xlsx", ".html", ".htm", ".md", ".txt"}
MOTS_PAR_SECTION = 600


def slug(nom: str) -> str:
    nom = unicodedata.normalize("NFD", nom.lower())
    nom = "".join(c for c in nom if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-z0-9]+", "-", nom).strip("-")[:80] or "document"


def _decouper_texte(texte: str, prefixe: str) -> str:
    """Regroupe un texte sans titres en sections « ## » d'environ MOTS_PAR_SECTION mots."""
    paragraphes = [p.strip() for p in re.split(r"\n\s*\n", texte) if p.strip()]
    sections, courant, n = [], [], 0
    for p in paragraphes:
        courant.append(p)
        n += len(p.split())
        if n >= MOTS_PAR_SECTION:
            sections.append(courant)
            courant, n = [], 0
    if courant:
        sections.append(courant)
    return "\n\n".join(
        f"## {prefixe} ({i}/{len(sections)})\n\n" + "\n\n".join(s) for i, s in enumerate(sections, 1)
    )


def convertir_pdf(chemin: Path) -> str:
    from pypdf import PdfReader

    pages = []
    for i, page in enumerate(PdfReader(chemin).pages, 1):
        texte = re.sub(r"[ \t]+", " ", page.extract_text() or "").strip()
        if texte:
            pages.append(f"## {chemin.stem} — page {i}\n\n{texte}")
    return "\n\n".join(pages)


def convertir_docx(chemin: Path) -> str:
    import docx

    document = docx.Document(chemin)
    lignes: list[str] = []
    for bloc in document.element.body.iterchildren():
        tag = bloc.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = docx.text.paragraph.Paragraph(bloc, document)
            texte = p.text.strip()
            if not texte:
                continue
            style = (p.style.name or "").lower() if p.style is not None else ""
            if style.startswith(("heading 1", "titre 1", "title", "titre")):
                lignes.append(f"\n## {texte}\n")
            elif style.startswith(("heading", "titre")):
                lignes.append(f"\n### {texte}\n")
            elif "list" in style or "liste" in style:
                lignes.append(f"- {texte}")
            else:
                lignes.append(f"\n{texte}\n")
        elif tag == "tbl":
            table = docx.table.Table(bloc, document)
            lignes.append("\n" + _table_markdown([[c.text.strip() for c in r.cells] for r in table.rows]) + "\n")
    texte = "\n".join(lignes)
    return texte if "\n## " in texte else _decouper_texte(texte, chemin.stem)


def convertir_xlsx(chemin: Path) -> str:
    from openpyxl import load_workbook

    classeur = load_workbook(chemin, data_only=True, read_only=True)
    sections = []
    for feuille in classeur.worksheets:
        lignes = [
            ["" if v is None else str(v).strip() for v in ligne]
            for ligne in feuille.iter_rows(values_only=True)
        ]
        lignes = [l for l in lignes if any(l)]
        if not lignes:
            continue
        # Tableaux découpés par blocs de 40 lignes, en répétant l'en-tête.
        entete, corps = lignes[0], lignes[1:] or [[""] * len(lignes[0])]
        for i in range(0, len(corps), 40):
            bloc = corps[i : i + 40]
            sections.append(f"## {feuille.title} (lignes {i + 2}–{i + 1 + len(bloc)})\n\n" + _table_markdown([entete, *bloc]))
    return "\n\n".join(sections)


def convertir_html(chemin: Path) -> str:
    from bs4 import BeautifulSoup

    soupe = BeautifulSoup(chemin.read_text(encoding="utf-8", errors="ignore"), "html.parser")
    for inutile in soupe(["script", "style", "noscript", "svg", "img", "nav", "footer"]):
        inutile.decompose()
    lignes: list[str] = []
    for el in soupe.find_all(["h1", "h2", "h3", "h4", "p", "li", "table"]):
        if el.find_parent("table") and el.name != "table":
            continue
        if el.name == "table":
            lignes.append(_table_markdown([[c.get_text(" ", strip=True) for c in tr.find_all(["td", "th"])] for tr in el.find_all("tr")]))
            continue
        texte = el.get_text(" ", strip=True)
        if not texte:
            continue
        if el.name in ("h1", "h2"):
            lignes.append(f"\n## {texte}\n")
        elif el.name in ("h3", "h4"):
            lignes.append(f"\n### {texte}\n")
        elif el.name == "li":
            lignes.append(f"- {texte}")
        else:
            lignes.append(f"\n{texte}\n")
    texte = "\n".join(lignes)
    return texte if "\n## " in texte else _decouper_texte(texte, chemin.stem)


def _table_markdown(lignes: list[list[str]]) -> str:
    lignes = [l for l in lignes if any(c for c in l)]
    if not lignes:
        return ""
    largeur = max(len(l) for l in lignes)
    lignes = [[c.replace("|", "/").replace("\n", " ") for c in l] + [""] * (largeur - len(l)) for l in lignes]
    rendu = ["| " + " | ".join(lignes[0]) + " |", "|" + "---|" * largeur]
    rendu += ["| " + " | ".join(l) + " |" for l in lignes[1:]]
    return "\n".join(rendu)


CONVERTISSEURS = {
    ".pdf": convertir_pdf,
    ".docx": convertir_docx,
    ".xlsx": convertir_xlsx,
    ".html": convertir_html,
    ".htm": convertir_html,
}


def importer(chemin: Path, acces: str, categorie: str, public: list[str]) -> Path:
    extension = chemin.suffix.lower()
    if extension in CONVERTISSEURS:
        corps = CONVERTISSEURS[extension](chemin)
    else:
        corps = chemin.read_text(encoding="utf-8", errors="ignore")
        if "\n## " not in f"\n{corps}":
            corps = _decouper_texte(corps, chemin.stem)
    if not corps.strip():
        raise ValueError("aucun texte extrait (document scanné ? passez-le d'abord dans un OCR)")

    titre = chemin.stem.replace("_", " ").replace("-", " ").strip()
    entete = (
        "---\n"
        f'titre: "{titre}"\n'
        f"categorie: {categorie}\n"
        f"public: [{', '.join(public)}]\n"
        f"acces: {acces}\n"
        f'source: "{chemin.name}"\n'
        f"mise_a_jour: {date.today():%Y-%m}\n"
        "---\n"
    )
    cible = DOSSIER_PRIVE / f"{slug(chemin.stem)}.md"
    cible.write_text(f"{entete}# {titre}\n\n{corps.strip()}\n", encoding="utf-8")
    return cible


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("fichiers", nargs="*", type=Path, help=f"par défaut : tout {DOSSIER_SOURCES}")
    parser.add_argument("--acces", choices=["interne", "public"], default="interne")
    parser.add_argument("--categorie", default="produit")
    parser.add_argument("--public", default="particulier,artisan,prescripteur,commercial", help="profils visés")
    args = parser.parse_args(argv)

    fichiers = args.fichiers or sorted(p for p in DOSSIER_SOURCES.glob("**/*") if p.suffix.lower() in FORMATS)
    if not fichiers:
        print(f"Aucun document trouvé. Déposez vos fichiers dans {DOSSIER_SOURCES}")
        return 1
    DOSSIER_PRIVE.mkdir(parents=True, exist_ok=True)
    public = [p.strip() for p in args.public.split(",") if p.strip()]
    echecs = 0
    for f in fichiers:
        try:
            cible = importer(f, args.acces, args.categorie, public)
            print(f"OK   {f.name} -> {cible.relative_to(RACINE)} ({args.acces})")
        except Exception as exc:  # un document illisible ne doit pas bloquer les autres
            echecs += 1
            print(f"ÉCHEC {f.name} : {exc}", file=sys.stderr)
    return 1 if echecs else 0


if __name__ == "__main__":
    raise SystemExit(main())
